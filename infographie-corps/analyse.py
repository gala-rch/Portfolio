"""
Analyse des données de Wee, Sznycer & Krems (2025), Science Advances
« Ce que vaut notre corps » — refonte de l'infographie Epsiloon.

Ce script :
  1. lit donnees/Data.xlsx (feuille « US data ») ;
  2. garde les participants retenus par les auteurs (Data inclusion = 1) ;
  3. calcule, pour chaque critère et chaque partie du corps :
     moyenne, médiane, écart-type, n, et le rang (1 = la plus précieuse) ;
  4. repère les parties qui changent le plus de rang d'un critère à l'autre ;
  5. compare hommes et femmes ;
  6. estime la stabilité des rangs (bootstrap) ;
  7. écrit tout dans sorties/resultats_USA.xlsx + des CSV.

Lancer :  python3 analyse.py
Pour changer un réglage, modifier la section « RÉGLAGES » ci-dessous.
"""

from pathlib import Path

import numpy as np
import pandas as pd

ICI = Path(__file__).parent
SORTIES = ICI / "sorties"
SORTIES.mkdir(exist_ok=True)

# ---------------------------------------------------------------- RÉGLAGES
FICHIER = ICI / "donnees" / "Data.xlsx"
FEUILLE = "US data"
FILTRER = True            # True = seulement Data inclusion = 1 (méthode des auteurs)
N_BOOTSTRAP = 2000        # tirages pour estimer la stabilité des rangs
GRAINE = 42               # pour que le bootstrap donne toujours le même résultat

# Les 5 critères de l'affiche, dans l'ordre des colonnes (gauche → droite)
CRITERES = {
    "Difficulty":   "Difficulté à s'en passer",
    "My Anger":     "Colère si on en perd l'usage",
    "My Gratitude": "Gratitude si on en retrouve l'usage",
    "Price":        "Prix sur le marché",
    "Compensation": "Dédommagement à verser",
}
# Les 2 critères de l'étude absents de l'affiche (calculés quand même, à part)
CRITERES_BONUS = {
    "Their Anger":     "Colère supposée d'autrui",
    "Their Gratitude": "Gratitude supposée d'autrui",
}

# Colonne du fichier → (libellé FR de l'affiche, groupe, ordre sur l'affiche)
# Les 4 colonnes absentes de l'affiche (testicules, une oreille, un œil)
# ne sont pas listées ici, donc ignorées.
PARTIES = {
    "Total paralysis":                       ("Motricité globale",     "Fonctions & sens", 1),
    "Imbecility or total mental deficiency": ("Capacités mentales",    "Fonctions & sens", 2),
    "The sight of both eyes":                ("Vue des 2 yeux",        "Fonctions & sens", 3),
    "Both arms":                             ("Les 2 bras",            "Bras & mains", 4),
    "Both legs":                             ("Les 2 jambes",          "Jambes & pieds", 5),
    "Both hands":                            ("Les 2 mains",           "Bras & mains", 6),
    "The ability to speak and chew":         ("Parole et mastication", "Fonctions & sens", 7),
    "Both feet":                             ("Les 2 pieds",           "Jambes & pieds", 8),
    "The hearing in both ears":              ("Ouïe des 2 oreilles",   "Fonctions & sens", 9),
    "All ten fingers":                       ("Les 10 doigts",         "Bras & mains", 10),
    "One arm":                               ("Bras",                  "Bras & mains", 11),
    "One leg":                               ("Jambe",                 "Jambes & pieds", 12),
    "Total disfigurement of face":           ("Visage",                "Visage, dents & intimité", 13),
    "One hand":                              ("Main",                  "Bras & mains", 14),
    "Thumb":                                 ("Pouce",                 "Bras & mains", 15),
    "One foot":                              ("Pied",                  "Jambes & pieds", 16),
    "The nose":                              ("Nez",                   "Visage, dents & intimité", 17),
    "All ten toes":                          ("Les 10 orteils",        "Jambes & pieds", 18),
    "Genitals":                              ("Parties génitales",     "Visage, dents & intimité", 19),
    "The sight of one eye":                  ("Vue d'un œil",          "Fonctions & sens", 20),
    "Index finger":                          ("Index",                 "Bras & mains", 21),
    "The hearing in one ear":                ("Ouïe d'une oreille",    "Fonctions & sens", 22),
    "Middle finger":                         ("Majeur",                "Bras & mains", 23),
    "Big toe":                               ("Gros orteil",           "Jambes & pieds", 24),
    "Ring finger":                           ("Annulaire",             "Bras & mains", 25),
    "Little finger":                         ("Auriculaire",           "Bras & mains", 26),
    "Second toe":                            ("2e orteil",             "Jambes & pieds", 27),
    "Third toe":                             ("3e orteil",             "Jambes & pieds", 28),
    "One front tooth":                       ("Incisive",              "Visage, dents & intimité", 29),
    "Fourth toe":                            ("4e orteil",             "Jambes & pieds", 30),
    "Fifth or little toe":                   ("Petit orteil",          "Jambes & pieds", 31),
    "One canine tooth":                      ("Canine",                "Visage, dents & intimité", 32),
    "One molar tooth":                       ("Molaire",               "Visage, dents & intimité", 33),
    "One big toenail":                       ("Ongle de pied",         "Jambes & pieds", 34),
    "One fingernail":                        ("Ongle de main",         "Bras & mains", 35),
}
GROUPES = ["Fonctions & sens", "Bras & mains", "Jambes & pieds", "Visage, dents & intimité"]


# ---------------------------------------------------------------- 1. LECTURE
def lire():
    df = pd.read_excel(FICHIER, sheet_name=FEUILLE)
    # Nettoyage des en-têtes : espaces en trop ("The sight of both eyes ",
    # "Middle  finger") et intitulés longs raccourcis.
    df.columns = [" ".join(str(c).split()) for c in df.columns]
    renommer = {}
    for c in df.columns:
        if c.startswith("Data inclusion"):
            renommer[c] = "Inclusion"
        elif c.startswith("Attention check"):
            renommer[c] = "Attention"
        elif c.startswith("Highest level of schooling"):
            renommer[c] = "Etudes"
        elif c.startswith("Major area of study"):
            renommer[c] = "Domaine"
    return df.rename(columns=renommer)


def filtrer(df):
    return df[df["Inclusion"] == 1].copy() if FILTRER else df.copy()


# ---------------------------------------------------------------- 2. CALCULS
def table_longue(d, criteres):
    """Une ligne par (critère, partie) : stats + rang parmi les 35 parties."""
    lignes = []
    for cond, g in d.groupby("Condition"):
        if cond not in criteres:
            continue
        for col in PARTIES:
            v = g[col].dropna()          # cellule vide = pas de réponse
            lignes.append({
                "critere_en": cond,
                "critere": criteres[cond],
                "partie_en": col,
                "partie": PARTIES[col][0],
                "groupe": PARTIES[col][1],
                "ordre_affiche": PARTIES[col][2],
                "moyenne": v.mean(),
                "mediane": v.median(),
                "ecart_type": v.std(),
                "n": len(v),
            })
    t = pd.DataFrame(lignes)
    # Rang : 1 = moyenne la plus élevée = la plus précieuse.
    # On compare des rangs (pas des valeurs) car Prix et Dédommagement sont
    # notés de 0 à 10 alors que les autres critères vont de 0 à 100.
    t["rang"] = (t.groupby("critere_en")["moyenne"]
                  .rank(ascending=False, method="first").astype(int))
    ordre = {c: i for i, c in enumerate(criteres)}
    return t.sort_values(["critere_en", "rang"], key=lambda s: s.map(ordre) if s.name == "critere_en" else s)


def large(t, valeur, criteres):
    """Parties en lignes, critères en colonnes."""
    w = t.pivot(index="partie", columns="critere", values=valeur)
    w = w[[criteres[c] for c in criteres]]
    info = t.drop_duplicates("partie").set_index("partie")[["groupe", "ordre_affiche"]]
    w = info.join(w).sort_values("ordre_affiche")
    return w


def mouvements(rangs):
    """Mesure à quel point chaque partie change de rang d'un critère à l'autre."""
    r = rangs[[CRITERES[c] for c in CRITERES]]
    pas = r.diff(axis=1).iloc[:, 1:]            # variation entre colonnes voisines
    noms_pas = [f"{a} → {b}" for a, b in zip(r.columns[:-1], r.columns[1:])]
    pas.columns = noms_pas
    m = pd.DataFrame(index=r.index)
    m["groupe"] = rangs["groupe"]
    m["rang_min"] = r.min(axis=1)
    m["rang_max"] = r.max(axis=1)
    m["amplitude"] = m["rang_max"] - m["rang_min"]
    m["cumul_des_sauts"] = pas.abs().sum(axis=1)  # distance parcourue par le ruban
    idx = pas.abs().values.argmax(axis=1)
    m["plus_gros_saut"] = [pas.iloc[i, j] for i, j in enumerate(idx)]
    m["ou"] = [noms_pas[j] for j in idx]
    m["difficulte_vers_dedommagement"] = r.iloc[:, -1] - r.iloc[:, 0]
    m = m.join(pas)
    return m.sort_values("cumul_des_sauts", ascending=False)


def par_sexe(d):
    """Rangs calculés séparément chez les hommes et chez les femmes."""
    res = []
    for sexe, g in d.groupby("Sex"):
        t = table_longue(g, CRITERES)
        t["sexe"] = sexe
        res.append(t)
    t = pd.concat(res)
    rh = large(t[t.sexe == "Male"], "rang", CRITERES)
    rf = large(t[t.sexe == "Female"], "rang", CRITERES)
    cols = [CRITERES[c] for c in CRITERES]
    out = rh[["groupe", "ordre_affiche"]].copy()
    out["rang_moyen_hommes"] = rh[cols].mean(axis=1)
    out["rang_moyen_femmes"] = rf[cols].mean(axis=1)
    # négatif = les hommes la classent plus haut (plus précieuse) que les femmes
    out["ecart_H_moins_F"] = out["rang_moyen_hommes"] - out["rang_moyen_femmes"]
    for c in cols:
        out[f"H−F · {c}"] = rh[c] - rf[c]
    n = d.groupby(["Condition", "Sex"]).size().unstack()
    return out.sort_values("ecart_H_moins_F"), n


def bootstrap_rangs(d):
    """Intervalle à 95 % du rang : on retire au hasard (avec remise) les
    participants de chaque critère, on recalcule les rangs, 2000 fois."""
    rng = np.random.default_rng(GRAINE)
    cols = list(PARTIES)
    res = []
    for cond, lib in CRITERES.items():
        x = d.loc[d.Condition == cond, cols].to_numpy(dtype=float)
        tirages = np.empty((N_BOOTSTRAP, len(cols)))
        for b in range(N_BOOTSTRAP):
            ech = x[rng.integers(0, len(x), len(x))]
            moy = np.nanmean(ech, axis=0)
            tirages[b] = pd.Series(-moy).rank(method="first").to_numpy()
        for j, col in enumerate(cols):
            res.append({"critere": lib, "partie": PARTIES[col][0],
                        "rang_bas_2.5%": np.percentile(tirages[:, j], 2.5),
                        "rang_haut_97.5%": np.percentile(tirages[:, j], 97.5)})
    return pd.DataFrame(res)


# ---------------------------------------------------------------- 3. EXPORT
def exporter(brut, d, t, t_bonus, rangs, moyennes, mouv, sexe, n_sexe, boot, comp):
    csv = dict(sep=";", decimal=",", encoding="utf-8-sig", index=False)  # Excel FR
    t.to_csv(SORTIES / "long_USA.csv", **csv)
    rangs.reset_index().to_csv(SORTIES / "rangs_USA.csv", **csv)
    moyennes.reset_index().to_csv(SORTIES / "moyennes_USA.csv", **csv)

    lisez_moi = pd.DataFrame({"Onglet": [
        "rangs", "moyennes", "long", "mouvements", "hommes_femmes", "effectifs",
        "stabilite_rangs", "filtre_ou_non", "criteres_bonus", "correspondance",
        "donnees_filtrees"], "Contenu": [
        "Rang de chaque partie (1 = la plus précieuse) pour les 5 critères de l'affiche. C'est la table qui sert à dessiner les rubans.",
        "Moyenne des notes. Attention : Prix et Dédommagement vont de 0 à 10, les 3 autres de 0 à 100.",
        "Format long : une ligne par critère × partie avec moyenne, médiane, écart-type, n et rang.",
        "Les parties qui bougent le plus. cumul_des_sauts = distance totale parcourue par le ruban ; amplitude = rang max − rang min.",
        "Rang moyen (sur les 5 critères) chez les hommes et chez les femmes. Écart négatif = partie plus précieuse pour les hommes. ~20 personnes par case : prudence.",
        "Nombre de participants par critère et par sexe après filtrage.",
        "Intervalle à 95 % du rang par bootstrap. Si deux intervalles se chevauchent, l'ordre entre ces deux parties n'est pas solide.",
        "Rangs avec et sans le filtre Data inclusion. L'affiche dit « 306 Américains » : elle a sans doute utilisé les 306 sans filtre.",
        "Les 2 critères de l'étude non montrés sur l'affiche (colère / gratitude supposées d'autrui).",
        "Correspondance colonnes du fichier ↔ libellés de l'affiche ↔ groupe.",
        "Les données individuelles gardées pour l'analyse (après filtre).",
    ]})
    corresp = pd.DataFrame([(k, *v) for k, v in PARTIES.items()],
                           columns=["colonne_fichier", "libelle_affiche", "groupe", "ordre_affiche"])
    cols_d = ["Participant #", "Condition", "Sex", "Age"] + list(PARTIES)

    with pd.ExcelWriter(SORTIES / "resultats_USA.xlsx", engine="openpyxl") as xw:
        lisez_moi.to_excel(xw, sheet_name="LISEZ-MOI", index=False)
        rangs.to_excel(xw, sheet_name="rangs")
        moyennes.round(2).to_excel(xw, sheet_name="moyennes")
        t.round(2).to_excel(xw, sheet_name="long", index=False)
        mouv.to_excel(xw, sheet_name="mouvements")
        sexe.round(2).to_excel(xw, sheet_name="hommes_femmes")
        n_sexe.to_excel(xw, sheet_name="effectifs")
        boot.to_excel(xw, sheet_name="stabilite_rangs", index=False)
        comp.to_excel(xw, sheet_name="filtre_ou_non")
        large(t_bonus, "rang", CRITERES_BONUS).to_excel(xw, sheet_name="criteres_bonus")
        corresp.to_excel(xw, sheet_name="correspondance", index=False)
        d[cols_d].to_excel(xw, sheet_name="donnees_filtrees", index=False)
        mettre_en_forme(xw)


def mettre_en_forme(xw):
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    for ws in xw.book.worksheets:
        ws.freeze_panes = "B2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EEEAE0")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 45
        for col in ws.columns:
            largeur = max(len(str(c.value or "")) for c in list(col)[1:40] + [col[0]])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largeur + 2), 60)
    xw.book["LISEZ-MOI"].column_dimensions["B"].width = 120
    # Rangs colorés : foncé = précieux, clair = peu précieux
    ws = xw.book["rangs"]
    ws.conditional_formatting.add(f"D2:H{ws.max_row}", ColorScaleRule(
        start_type="num", start_value=1, start_color="1C5CAB",
        mid_type="num", mid_value=18, mid_color="9EC5F4",
        end_type="num", end_value=35, end_color="FFFFFF"))


# ---------------------------------------------------------------- MAIN
def main():
    brut = lire()
    d = filtrer(brut)
    print(f"USA : {len(brut)} participants, {len(d)} gardés après filtre.")
    print(d.groupby("Condition").size().to_string(), "\n")

    t = table_longue(d, CRITERES)
    t_bonus = table_longue(d, CRITERES_BONUS)
    rangs = large(t, "rang", CRITERES)
    moyennes = large(t, "moyenne", CRITERES)
    mouv = mouvements(rangs)
    sexe, n_sexe = par_sexe(d)
    boot = bootstrap_rangs(d)

    # Comparaison avec / sans filtre
    r_sans = large(table_longue(brut, CRITERES), "rang", CRITERES)
    comp = rangs[["groupe"]].copy()
    for c in CRITERES.values():
        comp[f"{c} (filtré, n≈40)"] = rangs[c]
        comp[f"{c} (sans filtre)"] = r_sans[c]

    exporter(brut, d, t, t_bonus, rangs, moyennes, mouv, sexe, n_sexe, boot, comp)

    print("Rangs (1 = la plus précieuse) :")
    print(rangs.drop(columns="ordre_affiche").to_string(), "\n")
    print("Parties qui bougent le plus :")
    print(mouv[["groupe", "rang_min", "rang_max", "amplitude", "cumul_des_sauts",
                "plus_gros_saut", "ou"]].head(12).to_string(), "\n")
    print("Écarts hommes / femmes (rang moyen) :")
    print(pd.concat([sexe.head(6), sexe.tail(6)])[
        ["rang_moyen_hommes", "rang_moyen_femmes", "ecart_H_moins_F"]].round(1).to_string())
    print(f"\nFichiers écrits dans {SORTIES}")


if __name__ == "__main__":
    main()
